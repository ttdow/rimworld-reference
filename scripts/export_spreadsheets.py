"""Export the source workbooks and CSVs to the dependency-free reference website.

Read-only access to the originals. Requires openpyxl for extraction only.
Run from any directory: python scripts/export_spreadsheets.py
"""

from __future__ import annotations

import csv
import hashlib
import json
import math
import shutil
from pathlib import Path

import openpyxl


ROOT = Path(__file__).resolve().parents[1]
SITE = ROOT / "site"

SCHEMAS = {
    "apparel": {
        "file": "Apparel.xlsx",
        "name": "Apparel",
        "nameKey": "name",
        "categoryKey": "layer",
        "columns": [
            ("layer", "Layer", "text"),
            ("name", "Name", "text"),
            ("sharp", "Armor - Sharp", "number"),
            ("blunt", "Armor - Blunt", "number"),
            ("heatArmor", "Armor - Heat", "number"),
            ("hp", "HP (Hit Points)", "number"),
            ("coverage", "Coverage", "text"),
            ("coldMultiplier", "Cold Insulation Multiplier", "multiplier"),
            ("coldBase", "Cold Insulation Base (°C)", "temperature"),
            ("heatMultiplier", "Heat Insulation Multiplier", "multiplier"),
            ("heatBase", "Heat Insulation Base (°C)", "temperature"),
            ("special", "Special Properties", "text"),
            ("research", "Research requirement", "text"),
            ("work", "Work to Make (ticks)", "number"),
            ("material1Amount", "Material 1 Amount", "number"),
            ("material1Type", "Material 1 Type", "text"),
            ("material2Amount", "Material 2 Amount", "number"),
            ("material2Type", "Material 2 Type", "text"),
            ("marketValue", "Base Market Value", "number"),
            ("valueWork", "Value/Work", "number"),
            ("valueMaterial", "Value/Material", "number"),
        ],
    },
    "textiles": {
        "file": "Textiles.xlsx",
        "name": "Textiles",
        "nameKey": "name",
        "categoryKey": "category",
        "columns": [
            ("name", "Textile", "text"),
            ("category", "Stuff category", "text"),
            ("beauty", "Beauty multiplier", "multiplier"),
            ("hp", "Max HP multiplier", "multiplier"),
            ("flammability", "Flammability multiplier", "multiplier"),
            ("sharp", "Armor sharp multiplier", "multiplier"),
            ("blunt", "Armor blunt multiplier", "multiplier"),
            ("heatArmor", "Armor heat multiplier", "multiplier"),
            ("coldBonus", "Cold insulation material bonus (C)", "temperature"),
            ("heatBonus", "Heat insulation material bonus (C)", "temperature"),
            ("marketValue", "Market value (silver/unit)", "number"),
        ],
    },
}

EGG_SCHEMA = {
    "file": "rimworld_egg_laying_animals_comparison.csv",
    "name": "Egg-laying animals",
    "nameKey": "name",
    "categoryKey": "diet",
    "columns": [
        ("name", "Animal", "text"),
        ("content", "Content", "text"),
        ("diet", "Diet", "text"),
        ("unfertilized", "Lays unfertilized eggs (no male)", "text"),
        ("maleRequired", "Male required for egg output", "text"),
        ("eggNutrition", "Egg nutrition (per egg)", "number"),
        ("clutch", "Average eggs per clutch", "number"),
        ("intervalDays", "Egg laying interval (days)", "number"),
        ("eggsPerDay", "Average eggs per day (fertilized if required)", "number"),
        ("nutritionPerDay", "Egg nutrition per day (female, max)", "number"),
        ("femaleFeed", "Adult female daily feed nutrition", "number"),
        ("femaleEfficiency", "Gross egg nutrition / female feed (%)", "percent"),
        ("malesPerFemale", "Approx males per female for continuous fertilization", "number"),
        ("breederEfficiency", "Egg nutrition / total breeder feed (%)", "percent"),
        ("annualEggs", "Expected eggs per 60-day year (max)", "number"),
        ("wikiUrl", "Species wiki URL", "url"),
        ("sourceUrl", "Data source URL", "url"),
    ],
}


MILK_SCHEMA = {
    "file": "rimworld_milk_producing_animals_comparison.csv",
    "name": "Milk-producing animals",
    "nameKey": "name",
    "categoryKey": "packAnimal",
    "columns": [
        ("name", "Animal", "text"),
        ("content", "Content", "text"),
        ("diet", "Diet", "text"),
        ("adultFemales", "Milk from adult females", "text"),
        ("milkMaleRequired", "Male required for milk", "text"),
        ("milkPerMilking", "Milk per milking (units)", "number"),
        ("milkingInterval", "Milking interval (days)", "number"),
        ("milkPerDay", "Average milk/day (units)", "number"),
        ("milkNutrition", "Milk nutrition (per unit)", "number"),
        ("milkNutritionPerDay", "Milk nutrition/day (female, max)", "number"),
        ("femaleFeed", "Adult female feed nutrition/day", "number"),
        ("milkFeedEfficiency", "Gross milk nutrition / female feed (%)", "percent"),
        ("milkUnitsPerFeed", "Milk units / feed nutrition", "number"),
        ("hayPerDay", "Hay equivalent/day (units)", "number"),
        ("hay30Days", "Hay required for 30 days (units)", "number"),
        ("annualMilk", "Average milk in 60-day year (units)", "number"),
        ("packAnimal", "Pack animal", "text"),
        ("packCapacity", "Caravan pack capacity (kg)", "number"),
        ("caravanMount", "Caravan mount", "text"),
        ("minTemp", "Minimum comfortable temp (C)", "temperature"),
        ("maxTemp", "Maximum comfortable temp (C)", "temperature"),
        ("handlingSkill", "Minimum handling skill", "number"),
        ("maturityDays", "Maturity age (days)", "number"),
        ("wildness", "Wildness (%)", "percent"),
        ("notes", "Notes", "text"),
        ("wikiUrl", "Species wiki URL", "url"),
        ("sourceUrl", "Data source URL", "url"),
    ],
}


def read_csv(schema: dict) -> dict:
    source = ROOT / schema["file"]
    with source.open(encoding="utf-8-sig", newline="") as csv_file:
        reader = csv.reader(csv_file, strict=True)
        if next(reader, None) != [column[1] for column in schema["columns"]]:
            raise ValueError(f"{source.name}: column headers changed; update its schema before exporting.")
        rows = []
        names = set()
        while True:
            source_row = reader.line_num + 1
            cells = next(reader, None)
            if cells is None:
                break
            if not cells or all(value == "" for value in cells):
                continue
            record = {}
            for column, value in zip(schema["columns"], cells, strict=True):
                if value == "":
                    value = None
                elif column[2] in ("number", "percent", "temperature"):
                    value = float(value)
                    if not math.isfinite(value):
                        raise ValueError(f"{source.name}: non-finite number on row {source_row}.")
                record[column[0]] = value
            if not record["name"] or record["name"] in names:
                raise ValueError(f"{source.name}: missing or duplicate animal on row {source_row}.")
            names.add(record["name"])
            record["sourceRow"] = source_row
            rows.append(record)
    return {
        "name": schema["name"],
        "nameKey": schema["nameKey"],
        "categoryKey": schema["categoryKey"],
        "source": {"file": source.name, "format": "csv", "sha256": hashlib.sha256(source.read_bytes()).hexdigest()},
        "columns": [{"key": col[0], "label": col[1], "type": col[2]} for col in schema["columns"]],
        "rows": rows,
    }


def export() -> None:
    datasets = {}
    for key, schema in SCHEMAS.items():
        source = ROOT / schema["file"]
        workbook = openpyxl.load_workbook(source, data_only=True)
        formulas = openpyxl.load_workbook(source, data_only=False)
        try:
            if len(workbook.worksheets) != 1:
                raise ValueError(f"{source.name}: expected one sheet; review new sheets before exporting.")
            sheet = workbook.worksheets[0]
            headers = [column[1] for column in schema["columns"]]
            values = list(sheet.iter_rows(values_only=True))
            if list(values[0]) != headers:
                raise ValueError(f"{source.name}: column headers changed; update SCHEMAS before exporting.")

            # Formula values must be cached by the spreadsheet application.
            for row in formulas.worksheets[0]:
                for cell in row:
                    if cell.data_type == "f" and sheet[cell.coordinate].value is None:
                        raise ValueError(f"{source.name}!{cell.coordinate}: missing formula result.")

            rows = []
            for source_row, cells in enumerate(values[1:], start=2):
                if all(value is None for value in cells):
                    continue
                record = {column[0]: value for column, value in zip(schema["columns"], cells, strict=True)}
                if not record[schema["nameKey"]]:
                    raise ValueError(f"{source.name}: row {source_row} has no item name.")
                record["sourceRow"] = source_row
                rows.append(record)
            names = [row[schema["nameKey"]] for row in rows]
            if len(names) != len(set(names)):
                raise ValueError(f"{source.name}: duplicate item names need review.")
            datasets[key] = {
                "name": schema["name"],
                "nameKey": schema["nameKey"],
                "categoryKey": schema["categoryKey"],
                "source": {
                    "file": source.name,
                    "sheet": sheet.title,
                    "sha256": hashlib.sha256(source.read_bytes()).hexdigest(),
                },
                "columns": [{"key": col[0], "label": col[1], "type": col[2]} for col in schema["columns"]],
                "rows": rows,
            }
        finally:
            workbook.close()
            formulas.close()

    datasets["eggs"] = read_csv(EGG_SCHEMA)
    datasets["milk"] = read_csv(MILK_SCHEMA)
    assets = SITE / "assets"
    sources = SITE / "sources"
    assets.mkdir(parents=True, exist_ok=True)
    sources.mkdir(parents=True, exist_ok=True)
    payload = json.dumps({"datasets": datasets}, ensure_ascii=True, indent=2, allow_nan=False)
    (assets / "data.js").write_text("// Generated by scripts/export_spreadsheets.py.\nwindow.RIMWORLD_DATA = " + payload + ";\n", encoding="utf-8")
    for schema in [*SCHEMAS.values(), EGG_SCHEMA, MILK_SCHEMA]:
        shutil.copyfile(ROOT / schema["file"], sources / schema["file"])
    print("Exported " + ", ".join(f"{len(data['rows'])} {key} rows / {len(data['columns'])} columns" for key, data in datasets.items()))


if __name__ == "__main__":
    export()
